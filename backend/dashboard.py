from datetime import datetime
from io import BytesIO
from typing import Any, Dict, List, Optional

import numpy as np
import pandas as pd
from fastapi import APIRouter, HTTPException
from fastapi.responses import StreamingResponse
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import cm
from reportlab.platypus import PageBreak, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

from database import get_connection


router = APIRouter(prefix="/api/dashboard", tags=["dashboard"])

HEADER_FILL = PatternFill(fill_type="solid", fgColor="27AE60")
HEADER_FONT = Font(bold=True, color="FFFFFF")
ALT_ROW_FILL = PatternFill(fill_type="solid", fgColor="F0FAF4")


def _to_float(value: Any, default: float = 0.0) -> float:
    if value is None:
        return default
    try:
        if pd.isna(value):
            return default
    except Exception:
        pass
    try:
        return float(value)
    except (TypeError, ValueError):
        return default


def _to_int(value: Any, default: int = 0) -> int:
    return int(round(_to_float(value, float(default))))


def _run_query(sql: str) -> pd.DataFrame:
    conn = get_connection()
    if not conn:
        raise HTTPException(status_code=500, detail="Database connection failed")

    try:
        return pd.read_sql(sql, conn)
    except Exception as exc:
        raise HTTPException(status_code=500, detail=f"Dashboard query failed: {exc}") from exc
    finally:
        conn.close()


def _df_records(df: pd.DataFrame) -> List[Dict[str, Any]]:
    if df.empty:
        return []

    sanitized = df.replace({np.nan: None})
    return sanitized.to_dict(orient="records")


def _ensure_recommendation_logs_table() -> None:
    conn = get_connection()
    if not conn:
        return

    sql = (
        "CREATE TABLE IF NOT EXISTS recommendation_logs ("
        "id INT AUTO_INCREMENT PRIMARY KEY,"
        "product_name VARCHAR(200),"
        "weight_grams FLOAT,"
        "fragility VARCHAR(20),"
        "top_material VARCHAR(200),"
        "top_tier VARCHAR(20),"
        "top_cost FLOAT,"
        "top_co2 FLOAT,"
        "top_rank_score FLOAT,"
        "created_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP"
        ")"
    )

    cursor = None
    try:
        cursor = conn.cursor()
        cursor.execute(sql)
        conn.commit()
    finally:
        if cursor is not None:
            cursor.close()
        conn.close()


def _get_summary_data() -> Dict[str, Any]:
    _ensure_recommendation_logs_table()

    totals_df = _run_query(
        "SELECT "
        "COUNT(*) AS total_materials, "
        "AVG(co2_emission_score) AS avg_co2_score, "
        "AVG(packaging_cost_inr) AS avg_cost_inr, "
        "AVG(biodegradability_score) AS avg_biodegradability, "
        "AVG(recyclability_percent) AS avg_recyclability "
        "FROM materials"
    )

    counts_df = _run_query(
        "SELECT "
        "SUM(CASE WHEN sustainability_tier = 'Excellent' THEN 1 ELSE 0 END) AS excellent_materials, "
        "SUM(CASE WHEN sustainability_tier = 'Good' THEN 1 ELSE 0 END) AS good_materials "
        "FROM materials"
    )

    best_eco_df = _run_query(
        "SELECT packaging_material, co2_emission_score "
        "FROM materials "
        "ORDER BY co2_emission_score ASC "
        "LIMIT 1"
    )

    top_suitability_df = _run_query(
        "SELECT packaging_material, material_suitability_score "
        "FROM materials "
        "ORDER BY material_suitability_score DESC "
        "LIMIT 1"
    )

    tier_stats_df = _run_query(
        "SELECT sustainability_tier, "
        "AVG(co2_emission_score) AS avg_co2, "
        "AVG(packaging_cost_inr) AS avg_cost "
        "FROM materials "
        "GROUP BY sustainability_tier"
    )

    totals_row = totals_df.iloc[0] if not totals_df.empty else {}
    counts_row = counts_df.iloc[0] if not counts_df.empty else {}

    total_materials = _to_int(totals_row.get("total_materials", 0)) if isinstance(totals_row, pd.Series) else 0
    excellent_materials = _to_int(counts_row.get("excellent_materials", 0)) if isinstance(counts_row, pd.Series) else 0
    good_materials = _to_int(counts_row.get("good_materials", 0)) if isinstance(counts_row, pd.Series) else 0

    tier_lookup = {
        str(row.get("sustainability_tier", "")).strip().lower(): {
            "avg_co2": _to_float(row.get("avg_co2", 0)),
            "avg_cost": _to_float(row.get("avg_cost", 0)),
        }
        for row in _df_records(tier_stats_df)
    }

    poor_avg_co2 = tier_lookup.get("poor", {}).get("avg_co2", 0.0)
    excellent_avg_co2 = tier_lookup.get("excellent", {}).get("avg_co2", 0.0)
    poor_avg_cost = tier_lookup.get("poor", {}).get("avg_cost", 0.0)
    excellent_avg_cost = tier_lookup.get("excellent", {}).get("avg_cost", 0.0)

    co2_reduction_percent = round(((poor_avg_co2 - excellent_avg_co2) / poor_avg_co2) * 100, 1) if poor_avg_co2 > 0 else 0.0
    cost_savings_percent = round(((poor_avg_cost - excellent_avg_cost) / poor_avg_cost) * 100, 1) if poor_avg_cost > 0 else 0.0

    eco_adoption_rate = round(((excellent_materials + good_materials) / total_materials) * 100, 1) if total_materials > 0 else 0.0

    best_eco_material = (
        str(best_eco_df.iloc[0].get("packaging_material", "N/A"))
        if not best_eco_df.empty
        else "N/A"
    )
    highest_suitability_material = (
        str(top_suitability_df.iloc[0].get("packaging_material", "N/A"))
        if not top_suitability_df.empty
        else "N/A"
    )

    return {
        "total_materials": total_materials,
        "avg_co2_score": round(_to_float(totals_row.get("avg_co2_score", 0)) if isinstance(totals_row, pd.Series) else 0, 2),
        "avg_cost_inr": round(_to_float(totals_row.get("avg_cost_inr", 0)) if isinstance(totals_row, pd.Series) else 0, 2),
        "avg_biodegradability": round(_to_float(totals_row.get("avg_biodegradability", 0)) if isinstance(totals_row, pd.Series) else 0, 2),
        "avg_recyclability": round(_to_float(totals_row.get("avg_recyclability", 0)) if isinstance(totals_row, pd.Series) else 0, 2),
        "co2_reduction_percent": round(co2_reduction_percent, 1),
        "cost_savings_percent": round(cost_savings_percent, 1),
        "eco_adoption_rate": round(eco_adoption_rate, 1),
        "excellent_materials": excellent_materials,
        "good_materials": good_materials,
        "best_eco_material": best_eco_material,
        "highest_suitability_material": highest_suitability_material,
    }


@router.get("/summary")
def get_dashboard_summary() -> Dict[str, Any]:
    return _get_summary_data()


@router.get("/co2-analysis")
def get_co2_analysis() -> Dict[str, Any]:
    co2_by_tier = _run_query(
        "SELECT sustainability_tier, "
        "AVG(co2_emission_score) AS avg_co2, "
        "MIN(co2_emission_score) AS min_co2, "
        "MAX(co2_emission_score) AS max_co2, "
        "COUNT(*) AS count "
        "FROM materials "
        "GROUP BY sustainability_tier "
        "ORDER BY avg_co2 ASC"
    )

    co2_by_material_type = _run_query(
        "SELECT material_type, "
        "AVG(co2_emission_score) AS avg_co2, "
        "COUNT(*) AS count "
        "FROM materials "
        "GROUP BY material_type "
        "ORDER BY avg_co2 ASC"
    )

    co2_distribution = _run_query(
        "SELECT "
        "CASE "
        "WHEN co2_emission_score < 20 THEN 'Very Low 0-20' "
        "WHEN co2_emission_score < 40 THEN 'Low 20-40' "
        "WHEN co2_emission_score < 60 THEN 'Medium 40-60' "
        "ELSE 'High 60+' "
        "END AS bucket, "
        "COUNT(*) AS count "
        "FROM materials "
        "GROUP BY bucket"
    )

    top_10_lowest = _run_query(
        "SELECT packaging_material, co2_emission_score, sustainability_tier, material_type "
        "FROM materials "
        "ORDER BY co2_emission_score ASC "
        "LIMIT 10"
    )

    co2_by_industry = _run_query(
        "SELECT industry, AVG(co2_emission_score) AS avg_co2 "
        "FROM materials "
        "GROUP BY industry "
        "ORDER BY avg_co2 ASC "
        "LIMIT 10"
    )

    return {
        "co2_by_tier": _df_records(co2_by_tier),
        "co2_by_material_type": _df_records(co2_by_material_type),
        "co2_distribution_buckets": _df_records(co2_distribution),
        "top_10_lowest_co2": _df_records(top_10_lowest),
        "co2_by_industry": _df_records(co2_by_industry),
    }


@router.get("/cost-analysis")
def get_cost_analysis() -> Dict[str, Any]:
    cost_by_tier = _run_query(
        "SELECT sustainability_tier, "
        "AVG(packaging_cost_inr) AS avg_cost, "
        "MIN(packaging_cost_inr) AS min_cost, "
        "MAX(packaging_cost_inr) AS max_cost "
        "FROM materials "
        "GROUP BY sustainability_tier"
    )

    cost_by_material_type = _run_query(
        "SELECT material_type, "
        "AVG(packaging_cost_inr) AS avg_cost, "
        "COUNT(*) AS count "
        "FROM materials "
        "GROUP BY material_type "
        "ORDER BY avg_cost ASC"
    )

    cost_distribution = _run_query(
        "SELECT "
        "CASE "
        "WHEN packaging_cost_inr < 10 THEN 'Budget under 10' "
        "WHEN packaging_cost_inr < 30 THEN 'Mid 10-30' "
        "WHEN packaging_cost_inr < 60 THEN 'Premium 30-60' "
        "ELSE 'Luxury 60+' "
        "END AS cost_range, "
        "COUNT(*) AS count "
        "FROM materials "
        "GROUP BY cost_range"
    )

    cost_vs_suitability = _run_query(
        "SELECT packaging_cost_inr, material_suitability_score, sustainability_tier, packaging_material "
        "FROM materials "
        "ORDER BY RAND() "
        "LIMIT 200"
    )

    top_10_cheapest_eco = _run_query(
        "SELECT packaging_material, packaging_cost_inr, sustainability_tier, material_suitability_score "
        "FROM materials "
        "WHERE sustainability_tier IN ('Excellent','Good') "
        "ORDER BY packaging_cost_inr ASC "
        "LIMIT 10"
    )

    return {
        "cost_by_tier": _df_records(cost_by_tier),
        "cost_by_material_type": _df_records(cost_by_material_type),
        "cost_distribution": _df_records(cost_distribution),
        "cost_vs_suitability": _df_records(cost_vs_suitability),
        "top_10_cheapest_eco": _df_records(top_10_cheapest_eco),
    }


@router.get("/material-trends")
def get_material_trends() -> Dict[str, Any]:
    tier_distribution = _run_query(
        "SELECT sustainability_tier, COUNT(*) AS count, "
        "ROUND(COUNT(*) * 100.0 / (SELECT COUNT(*) FROM materials), 1) AS percentage "
        "FROM materials "
        "GROUP BY sustainability_tier "
        "ORDER BY count DESC"
    )

    material_type_distribution = _run_query(
        "SELECT material_type, COUNT(*) AS count, "
        "AVG(material_suitability_score) AS avg_suitability, "
        "AVG(packaging_cost_inr) AS avg_cost, "
        "AVG(biodegradability_score) AS avg_bio, "
        "AVG(recyclability_percent) AS avg_recycle, "
        "AVG(co2_emission_score) AS avg_co2 "
        "FROM materials "
        "GROUP BY material_type "
        "ORDER BY count DESC"
    )

    strength_distribution = _run_query(
        "SELECT strength, COUNT(*) AS count, AVG(co2_emission_score) AS avg_co2 "
        "FROM materials "
        "GROUP BY strength"
    )

    top_10_sustainable = _run_query(
        "SELECT packaging_material, material_type, industry, material_suitability_score, "
        "sustainability_tier, co2_emission_score, packaging_cost_inr, biodegradability_score, recyclability_percent "
        "FROM materials "
        "ORDER BY material_suitability_score DESC "
        "LIMIT 10"
    )

    biodegradability_by_type = _run_query(
        "SELECT material_type, "
        "AVG(biodegradability_score) AS avg_bio, "
        "AVG(recyclability_percent) AS avg_recycle "
        "FROM materials "
        "GROUP BY material_type "
        "ORDER BY avg_bio DESC"
    )

    industry_sustainability = _run_query(
        "SELECT industry, "
        "AVG(material_suitability_score) AS avg_suitability, "
        "AVG(co2_emission_score) AS avg_co2, "
        "COUNT(*) AS material_count "
        "FROM materials "
        "GROUP BY industry "
        "ORDER BY avg_suitability DESC "
        "LIMIT 10"
    )

    return {
        "tier_distribution": _df_records(tier_distribution),
        "material_type_distribution": _df_records(material_type_distribution),
        "strength_distribution": _df_records(strength_distribution),
        "top_10_sustainable": _df_records(top_10_sustainable),
        "biodegradability_by_type": _df_records(biodegradability_by_type),
        "industry_sustainability": _df_records(industry_sustainability),
    }


def _base_styles() -> Dict[str, ParagraphStyle]:
    styles = getSampleStyleSheet()
    styles.add(
        ParagraphStyle(
            name="TitleCenter",
            parent=styles["Title"],
            alignment=1,
            textColor=colors.HexColor("#1A1A2E"),
        )
    )
    styles.add(
        ParagraphStyle(
            name="HeadingGreen",
            parent=styles["Heading1"],
            textColor=colors.HexColor("#27AE60"),
            spaceAfter=10,
        )
    )
    styles.add(
        ParagraphStyle(
            name="Body",
            parent=styles["BodyText"],
            leading=16,
            textColor=colors.HexColor("#2C3E50"),
        )
    )
    return styles


def _page_footer(canvas, doc) -> None:
    canvas.saveState()
    canvas.setFont("Helvetica", 9)
    canvas.setFillColor(colors.HexColor("#2C3E50"))
    canvas.drawRightString(A4[0] - 2 * cm, 1.2 * cm, f"Page {doc.page}")
    canvas.restoreState()


def _styled_table(data: List[List[Any]], col_widths: Optional[List[float]] = None, alternate_rows: bool = False) -> Table:
    table = Table(data, colWidths=col_widths, repeatRows=1)
    style_rules = [
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#27AE60")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("ALIGN", (0, 0), (-1, -1), "LEFT"),
        ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#D7DBDD")),
        ("FONTSIZE", (0, 0), (-1, -1), 9),
        ("BOTTOMPADDING", (0, 0), (-1, 0), 8),
        ("TOPPADDING", (0, 0), (-1, 0), 8),
    ]

    if alternate_rows and len(data) > 2:
        for row_index in range(1, len(data)):
            if row_index % 2 == 0:
                style_rules.append(
                    ("BACKGROUND", (0, row_index), (-1, row_index), colors.HexColor("#F0FAF4"))
                )

    table.setStyle(TableStyle(style_rules))
    return table


def _summary_for_reports() -> Dict[str, Any]:
    return _get_summary_data()


def _safe_top_material_names(limit: int = 3) -> List[str]:
    df = _run_query(
        "SELECT packaging_material "
        "FROM materials "
        "ORDER BY material_suitability_score DESC "
        f"LIMIT {int(limit)}"
    )
    return [str(x) for x in df.get("packaging_material", []).tolist()] if not df.empty else []


@router.get("/export/pdf")
def export_dashboard_pdf() -> StreamingResponse:
    summary = _summary_for_reports()
    co2_df = _run_query(
        "SELECT sustainability_tier, AVG(co2_emission_score) AS avg_co2, "
        "MIN(co2_emission_score) AS min_co2, MAX(co2_emission_score) AS max_co2, COUNT(*) AS count "
        "FROM materials GROUP BY sustainability_tier ORDER BY avg_co2 ASC"
    )
    cost_df = _run_query(
        "SELECT material_type, AVG(packaging_cost_inr) AS avg_cost, COUNT(*) AS count "
        "FROM materials GROUP BY material_type ORDER BY avg_cost ASC"
    )
    top10_df = _run_query(
        "SELECT packaging_material, sustainability_tier, material_suitability_score, "
        "co2_emission_score, packaging_cost_inr "
        "FROM materials ORDER BY material_suitability_score DESC LIMIT 10"
    )
    cheapest_eco_df = _run_query(
        "SELECT packaging_material, packaging_cost_inr, sustainability_tier "
        "FROM materials WHERE sustainability_tier IN ('Excellent', 'Good') "
        "ORDER BY packaging_cost_inr ASC LIMIT 5"
    )

    styles = _base_styles()
    buffer = BytesIO()
    doc = SimpleDocTemplate(
        buffer,
        pagesize=A4,
        leftMargin=1.8 * cm,
        rightMargin=1.8 * cm,
        topMargin=1.8 * cm,
        bottomMargin=1.8 * cm,
    )

    elements: List[Any] = []
    today = datetime.now().strftime("%Y-%m-%d")

    elements.append(Spacer(1, 1.5 * cm))
    elements.append(Paragraph("EcoPackAI", styles["TitleCenter"]))
    elements.append(Spacer(1, 0.2 * cm))
    elements.append(Paragraph("Sustainability Report", styles["HeadingGreen"]))
    elements.append(Paragraph("AI-Powered Packaging Material Analysis", styles["Body"]))
    elements.append(Spacer(1, 0.4 * cm))
    elements.append(Paragraph(f"Generated on: {today}", styles["Body"]))
    elements.append(Paragraph("Powered by AI and Machine Learning", styles["Body"]))
    elements.append(Paragraph("Generated by EcoPackAI BI Dashboard", styles["Body"]))
    elements.append(PageBreak())

    elements.append(Paragraph("Executive Summary", styles["HeadingGreen"]))
    summary_table = [
        ["Metric", "Value"],
        ["Total Materials", f"{summary['total_materials']:,}"],
        ["CO2 Reduction Potential", f"{summary['co2_reduction_percent']}%"],
        ["Cost Savings Potential", f"{summary['cost_savings_percent']}%"],
        ["Eco Adoption Rate", f"{summary['eco_adoption_rate']}%"],
        ["Average Biodegradability", f"{summary['avg_biodegradability']} / 100"],
        ["Average Recyclability", f"{summary['avg_recyclability']}%"],
    ]
    elements.append(_styled_table(summary_table, [8 * cm, 7 * cm]))
    elements.append(Spacer(1, 0.5 * cm))
    elements.append(
        Paragraph(
            "CO2 reduction potential represents the expected emission drop when shifting from Poor tier to Excellent tier packaging materials.",
            styles["Body"],
        )
    )
    elements.append(
        Paragraph(
            "Cost savings potential estimates the average packaging cost improvement from replacing low-performing materials with eco-optimized alternatives.",
            styles["Body"],
        )
    )
    elements.append(PageBreak())

    elements.append(Paragraph("CO2 Emission Analysis", styles["HeadingGreen"]))
    co2_table_data = [["Sustainability Tier", "Avg CO2", "Min CO2", "Max CO2", "Material Count"]]
    for row in _df_records(co2_df):
        co2_table_data.append([
            row.get("sustainability_tier", "N/A"),
            f"{_to_float(row.get('avg_co2')):.2f}",
            f"{_to_float(row.get('min_co2')):.2f}",
            f"{_to_float(row.get('max_co2')):.2f}",
            f"{_to_int(row.get('count'))}",
        ])
    elements.append(_styled_table(co2_table_data, [3.4 * cm, 2.4 * cm, 2.4 * cm, 2.4 * cm, 3 * cm]))
    elements.append(Spacer(1, 0.4 * cm))
    elements.append(
        Paragraph(
            f"Materials in the Excellent tier produce about {summary['co2_reduction_percent']}% less CO2 than Poor tier materials.",
            styles["Body"],
        )
    )

    top_low_co2 = _run_query(
        "SELECT packaging_material, co2_emission_score FROM materials ORDER BY co2_emission_score ASC LIMIT 5"
    )
    low_table = [["Top 5 Low CO2 Materials", "CO2 Score"]]
    for row in _df_records(top_low_co2):
        low_table.append([
            row.get("packaging_material", "N/A"),
            f"{_to_float(row.get('co2_emission_score')):.2f}",
        ])
    elements.append(Spacer(1, 0.3 * cm))
    elements.append(_styled_table(low_table, [10 * cm, 5 * cm]))
    elements.append(PageBreak())

    elements.append(Paragraph("Cost Savings Analysis", styles["HeadingGreen"]))
    cost_table = [["Material Type", "Avg Cost INR", "Material Count", "Eco Rating"]]
    for row in _df_records(cost_df):
        eco_rating = "Strong" if _to_float(row.get("avg_cost")) <= summary["avg_cost_inr"] else "Moderate"
        cost_table.append([
            row.get("material_type", "N/A"),
            f"{_to_float(row.get('avg_cost')):.2f}",
            f"{_to_int(row.get('count'))}",
            eco_rating,
        ])
    elements.append(_styled_table(cost_table, [5 * cm, 3 * cm, 3 * cm, 4 * cm]))
    elements.append(Spacer(1, 0.4 * cm))
    elements.append(
        Paragraph(
            f"Switching to eco-friendly materials can save up to {summary['cost_savings_percent']}% in packaging costs when benchmarked against Poor tier options.",
            styles["Body"],
        )
    )

    cheap_table = [["Top 5 Cheapest Eco Materials", "Tier", "Cost INR"]]
    for row in _df_records(cheapest_eco_df):
        cheap_table.append([
            row.get("packaging_material", "N/A"),
            row.get("sustainability_tier", "N/A"),
            f"{_to_float(row.get('packaging_cost_inr')):.2f}",
        ])
    elements.append(Spacer(1, 0.3 * cm))
    elements.append(_styled_table(cheap_table, [8 * cm, 3 * cm, 4 * cm]))
    elements.append(PageBreak())

    elements.append(Paragraph("Top 10 Sustainable Materials", styles["HeadingGreen"]))
    top_table = [["Rank", "Material Name", "Tier", "Suitability", "CO2", "Cost INR"]]
    for idx, row in enumerate(_df_records(top10_df), start=1):
        top_table.append([
            idx,
            row.get("packaging_material", "N/A"),
            row.get("sustainability_tier", "N/A"),
            f"{_to_float(row.get('material_suitability_score')):.2f}",
            f"{_to_float(row.get('co2_emission_score')):.2f}",
            f"{_to_float(row.get('packaging_cost_inr')):.2f}",
        ])
    elements.append(_styled_table(top_table, [1.2 * cm, 6.0 * cm, 2.2 * cm, 2.3 * cm, 1.6 * cm, 2.3 * cm], alternate_rows=True))
    elements.append(PageBreak())

    elements.append(Paragraph("Recommendations and Action Plan", styles["HeadingGreen"]))
    top_three = _safe_top_material_names(3)
    recommendation_lines = [
        f"1. Switch to Excellent tier materials to reduce CO2 emissions by approximately {summary['co2_reduction_percent']}%.",
        "2. Prioritize Paper and Bioplastic material categories for strong sustainability-to-cost balance.",
        "3. Favor high biodegradability options to reduce landfill impact and improve end-of-life outcomes.",
        f"4. Consider top materials: {', '.join(top_three) if top_three else 'N/A'}.",
        "5. Build industry-specific procurement policies using top-ranked low-CO2, high-suitability materials.",
    ]
    for line in recommendation_lines:
        elements.append(Paragraph(line, styles["Body"]))
        elements.append(Spacer(1, 0.15 * cm))

    elements.append(Spacer(1, 0.3 * cm))
    elements.append(
        Paragraph(
            "Closing Note: This report is generated from live EcoPackAI material intelligence and is designed to support strategic sustainability decisions.",
            styles["Body"],
        )
    )

    doc.build(elements, onFirstPage=_page_footer, onLaterPages=_page_footer)
    buffer.seek(0)

    filename = f"EcoPackAI_Sustainability_Report_{today}.pdf"
    return StreamingResponse(
        buffer,
        media_type="application/pdf",
        headers={"Content-Disposition": f"attachment; filename={filename}"},
    )


def _style_sheet(ws) -> None:
    for row in ws.iter_rows(min_row=1, max_row=1):
        for cell in row:
            cell.fill = HEADER_FILL
            cell.font = HEADER_FONT
            cell.alignment = Alignment(horizontal="center", vertical="center")


def _auto_fit_columns(ws) -> None:
    for col in ws.columns:
        col_values = [str(cell.value) if cell.value is not None else "" for cell in col]
        max_len = max((len(v) for v in col_values), default=8)
        ws.column_dimensions[col[0].column_letter].width = min(max(max_len + 2, 10), 50)


def _write_df(ws, df: pd.DataFrame, start_row: int = 1, alternate_rows: bool = False) -> int:
    columns = list(df.columns)
    for col_index, col_name in enumerate(columns, start=1):
        ws.cell(row=start_row, column=col_index, value=col_name)

    for row_index, row in enumerate(df.itertuples(index=False), start=start_row + 1):
        for col_index, value in enumerate(row, start=1):
            ws.cell(row=row_index, column=col_index, value=value)

        if alternate_rows and row_index % 2 == 0:
            for col_index in range(1, len(columns) + 1):
                ws.cell(row=row_index, column=col_index).fill = ALT_ROW_FILL

    _style_sheet(ws)
    _auto_fit_columns(ws)
    return start_row + len(df) + 2


@router.get("/export/excel")
def export_dashboard_excel() -> StreamingResponse:
    summary = _summary_for_reports()

    co2_by_tier = _run_query(
        "SELECT sustainability_tier, AVG(co2_emission_score) AS avg_co2, "
        "MIN(co2_emission_score) AS min_co2, MAX(co2_emission_score) AS max_co2, COUNT(*) AS material_count "
        "FROM materials GROUP BY sustainability_tier ORDER BY avg_co2 ASC"
    )
    co2_by_type = _run_query(
        "SELECT material_type, AVG(co2_emission_score) AS avg_co2, COUNT(*) AS material_count "
        "FROM materials GROUP BY material_type ORDER BY avg_co2 ASC"
    )
    cost_by_tier = _run_query(
        "SELECT sustainability_tier, AVG(packaging_cost_inr) AS avg_cost, MIN(packaging_cost_inr) AS min_cost, "
        "MAX(packaging_cost_inr) AS max_cost FROM materials GROUP BY sustainability_tier"
    )
    cheapest_eco = _run_query(
        "SELECT packaging_material, packaging_cost_inr, sustainability_tier, material_suitability_score "
        "FROM materials WHERE sustainability_tier IN ('Excellent', 'Good') ORDER BY packaging_cost_inr ASC LIMIT 10"
    )
    top_materials = _run_query(
        "SELECT packaging_material, material_type, industry, sustainability_tier, material_suitability_score, "
        "co2_emission_score, packaging_cost_inr, biodegradability_score, recyclability_percent "
        "FROM materials ORDER BY material_suitability_score DESC LIMIT 10"
    )
    tier_distribution = _run_query(
        "SELECT sustainability_tier, COUNT(*) AS count, "
        "ROUND(COUNT(*) * 100.0 / (SELECT COUNT(*) FROM materials), 1) AS percentage "
        "FROM materials GROUP BY sustainability_tier ORDER BY count DESC"
    )
    material_types = _run_query(
        "SELECT material_type, COUNT(*) AS count, AVG(material_suitability_score) AS avg_suitability, "
        "AVG(packaging_cost_inr) AS avg_cost, AVG(co2_emission_score) AS avg_co2 "
        "FROM materials GROUP BY material_type ORDER BY count DESC"
    )
    industry_df = _run_query(
        "SELECT industry, AVG(material_suitability_score) AS avg_suitability, AVG(co2_emission_score) AS avg_co2, "
        "COUNT(*) AS material_count FROM materials GROUP BY industry ORDER BY avg_suitability DESC LIMIT 20"
    )
    raw_data = _run_query("SELECT * FROM materials LIMIT 500")

    wb = Workbook()

    ws_summary = wb.active
    ws_summary.title = "1_Summary"
    summary_df = pd.DataFrame(
        [
            ["Total Materials", summary["total_materials"]],
            ["CO2 Reduction Potential (%)", summary["co2_reduction_percent"]],
            ["Cost Savings Potential (%)", summary["cost_savings_percent"]],
            ["Eco Adoption Rate (%)", summary["eco_adoption_rate"]],
            ["Average Biodegradability", summary["avg_biodegradability"]],
            ["Average Recyclability", summary["avg_recyclability"]],
            ["Best Eco Material", summary["best_eco_material"]],
            ["Highest Suitability Material", summary["highest_suitability_material"]],
        ],
        columns=["Metric", "Value"],
    )
    _write_df(ws_summary, summary_df)

    ws_co2 = wb.create_sheet("2_CO2_Analysis")
    next_row = _write_df(ws_co2, co2_by_tier.rename(columns=str.title), start_row=1)
    next_row += 1
    ws_co2.cell(row=next_row, column=1, value="CO2 by Material Type")
    ws_co2.cell(row=next_row, column=1).font = Font(bold=True)
    _write_df(ws_co2, co2_by_type.rename(columns=str.title), start_row=next_row + 1)

    ws_cost = wb.create_sheet("3_Cost_Analysis")
    next_row = _write_df(ws_cost, cost_by_tier.rename(columns=str.title), start_row=1)
    next_row += 1
    ws_cost.cell(row=next_row, column=1, value="Top 10 Cheapest Eco Materials")
    ws_cost.cell(row=next_row, column=1).font = Font(bold=True)
    _write_df(ws_cost, cheapest_eco.rename(columns=str.title), start_row=next_row + 1)

    ws_top = wb.create_sheet("4_Top_Materials")
    top_df = top_materials.copy()
    top_df.insert(0, "rank", range(1, len(top_df) + 1))
    top_df = top_df.rename(
        columns={
            "rank": "Rank",
            "packaging_material": "Material Name",
            "material_type": "Type",
            "industry": "Industry",
            "sustainability_tier": "Tier",
            "material_suitability_score": "Suitability Score",
            "co2_emission_score": "CO2 Score",
            "packaging_cost_inr": "Cost INR",
            "biodegradability_score": "Biodegradability",
            "recyclability_percent": "Recyclability",
        }
    )
    _write_df(ws_top, top_df, alternate_rows=True)

    ws_tier = wb.create_sheet("5_Tier_Distribution")
    _write_df(ws_tier, tier_distribution.rename(columns=str.title))

    ws_types = wb.create_sheet("6_Material_Types")
    _write_df(ws_types, material_types.rename(columns=str.title))

    ws_industry = wb.create_sheet("7_Industry_Analysis")
    _write_df(ws_industry, industry_df.rename(columns=str.title))

    ws_raw = wb.create_sheet("8_Raw_Data")
    _write_df(ws_raw, raw_data)

    output = BytesIO()
    wb.save(output)
    output.seek(0)

    today = datetime.now().strftime("%Y-%m-%d")
    filename = f"EcoPackAI_Report_{today}.xlsx"
    return StreamingResponse(
        output,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename={filename}"},
    )
